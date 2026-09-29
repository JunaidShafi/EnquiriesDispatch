from playwright.sync_api import Playwright, sync_playwright
import mimetypes
import os
import smtplib
from email.utils import formataddr
from email.message import EmailMessage 
from dotenv import load_dotenv
from datetime import datetime
from email.utils import formataddr
import pandas as p
import openpyxl
from openpyxl.styles import Border, Font, PatternFill, Side
import re
load_dotenv()





def run(playwright: Playwright) -> None:
    browser = playwright.chromium.launch(headless=True)
    context = browser.new_context()
    page = context.new_page()
    page.goto("https://bomis.nascorptechnologies.com/Index")
    print(f"Navigating to {page.title()}")
    page.get_by_role("textbox", name="User Name").fill("junaidshafi@bomiskashmir.edu.in")
    page.get_by_role("textbox", name="Password").fill("Crap1234")
    page.get_by_role("button", name="Login").click()
    page.get_by_role("listitem").filter(has_text="Academic Setup Achievement/").get_by_role("link").click()
    page.get_by_role("link", name=" Admission Management").click()
    page.locator("a").filter(has_text="Reports").first.click()
    page.locator("a").filter(has_text="Dynamic Reports").click()
    page.wait_for_timeout(4000)
    page.goto("https://bomis.nascorptechnologies.com/gw/adm/dynamicAdmissionReportEdit?fl=aWQ9MTImX3BsXz1odHRwczovL2JvbWlzLm5hc2NvcnB0ZWNobm9sb2dpZXMuY29tL2d3L2Z3ay9hZG1fcmVwRHluYW1pYw==", wait_until="domcontentloaded")
    print(f"Navigating to {page.title()}")
    page.get_by_role("button", name="Click for Actions").click()
    with page.expect_download() as download_info:
        with page.expect_popup() as page1_info:
            page.get_by_role("link", name="Download Report (Excel)").click()
        page1 = page1_info.value
    download = download_info.value
    download.save_as("enq.xls")

    page1.close()

    # ---------------------


with sync_playwright() as playwright:
    run(playwright)


def cleanfile():
    print("Started Processing File")
    df = p.read_excel("enq.xls", header=4)
    df = df.dropna(subset=["Registration No."])
    df = df.iloc[:, 1:]
    df = df.fillna("")
    df.insert(0, "S No.", range(1, len(df) + 1))
    df = df.reset_index(drop=True)
    df = df.drop(columns="Form Sale Receipt No.")

    class_order = ["NURSERY", "KG1", "KG2"]

    df_sorted = df.copy()

    # FIX 1: Strip extra whitespace from the Class column so categories match perfectly
    df_sorted["Class"] = df_sorted["Class"].astype(str).str.strip()

    df_sorted["Class"] = p.Categorical(
        df_sorted["Class"], categories=class_order, ordered=True
    )

    # FIX 2: Sort by "Class" first, then by "Registration No."
    df_sorted = df_sorted.sort_values(
        by=["Class", "Registration No."]
    ).reset_index(drop=True)

    df_sorted["S No."] = range(1, len(df_sorted) + 1)
    file_path = "processed.xlsx"

    with p.ExcelWriter(file_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Enquiries")

        # FIX 3: Write df_sorted instead of df to the second sheet
        df_sorted.to_excel(writer, index=False, sheet_name="Sorted Classwise")

        def format_sheet(ws, dataframe):
            thin_border = Side(style="thin", color="000000")
            cell_border = Border(
                left=thin_border,
                right=thin_border,
                top=thin_border,
                bottom=thin_border,
            )

            font_header = Font(name="Calibri", size=11, bold=True)
            font_body = Font(name="Calibri", size=10, bold=False)
            header_fill = PatternFill(
                start_color="D3D3D3", end_color="D3D3D3", fill_type="solid"
            )

            # Format Header Row
            for cell in ws[1]:
                cell.font = font_header
                cell.border = cell_border
                cell.fill = header_fill

            # Format Data Rows
            for row in ws.iter_rows(min_row=2, max_row=len(dataframe) + 1):
                for cell in row:
                    cell.font = font_body
                    cell.border = cell_border

            # Autofit Columns
            for col in ws.columns:
                max_len = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        format_sheet(writer.sheets["Enquiries"], df)
        format_sheet(writer.sheets["Sorted Classwise"], df_sorted)


cleanfile()
def send_file(filepath):
        print("Started Sending Mail")
        SMTP_SERVER = 'smtp.gmail.com'
        SMTP_PORT = 465  # Use 465 for SSL or 587 for TLS
        GMAIL_USER = os.getenv("GMAIL_USER")
        GMAIL_APP_PASSWORD = os.getenv("GMAIL_APP_PASSWORD") # 16-character App Password

        msg = EmailMessage()
        msg['Subject'] = 'Daily Enquiry List'
        msg['From'] = formataddr(("Junaid Shafi",GMAIL_USER))
        msg['To'] = os.getenv("RECEIVER_EMAIL")
        msg.set_content(f'Registration details on {datetime.now().date()}')

        file_path = filepath
        mime_type, _ = mimetypes.guess_type(file_path)
        main_type, sub_type = (mime_type or 'application/octet-stream').split('/', 1)

        with open(file_path, 'rb') as f:
            msg.add_attachment(f.read(), maintype=main_type, subtype=sub_type, filename=os.path.basename(file_path))

        # 3. Connect to Gmail and send
        with smtplib.SMTP_SSL(SMTP_SERVER, SMTP_PORT) as server:
            server.login(GMAIL_USER, GMAIL_APP_PASSWORD)
            server.send_message(msg)
            print('Email successfully sent!')


send_file("processed.xlsx")
